#!/usr/bin/env python3
"""Render confirmed resume content into one of the Job Fit DOCX templates.

Usage:
    python3 render_resume.py resume.json --out 简历.docx [--pdf]

The input JSON holds only content the user has already confirmed; this script
never adds, rewrites or removes facts. Layout comes from the chosen template.
See references/resume-templates.md for the input format and the style rules.

Prints one JSON object to stdout, for example:
    {"ok": true, "docx": "...", "pdf": null, "pages": null, "warnings": [...]}

Requires python-docx (pip install python-docx). PDF export and page counting
need LibreOffice (soffice) on PATH; without it, only the DOCX is produced.
"""
from __future__ import annotations

import argparse
import json
import re
import shutil
import subprocess
import sys
import tempfile
from dataclasses import dataclass
from pathlib import Path

try:
    from docx import Document
    from docx.enum.section import WD_SECTION_START
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
except ImportError:  # pragma: no cover - reported to the caller as JSON
    print(json.dumps({"ok": False, "error": "python-docx is not installed; run: pip install python-docx"}, ensure_ascii=False))
    sys.exit(2)

FONT = "Arial"
FONT_EAST_ASIA = "Microsoft YaHei"
PAGE_WIDTH_CM = 21.0
TEMPLATES = ("professional-business", "technical-project", "one-page-compact", "ats-minimal")
LEGACY = {"ats-classic": "ats-minimal", "ats-compact": "professional-business", "ats-graduate": "technical-project"}
SECTION_TYPES = ("summary", "skills", "entries", "bullets")


# ---------- low-level formatting helpers ----------

def _ppr_child(paragraph, tag: str):
    ppr = paragraph._p.get_or_add_pPr()
    node = ppr.find(qn(tag))
    if node is None:
        node = OxmlElement(tag)
        ppr.append(node)
    return node


def shade(paragraph, fill: str) -> None:
    _ppr_child(paragraph, "w:shd").set(qn("w:fill"), fill)


def border(paragraph, side: str, color: str, size: int, space: int = 4) -> None:
    edge = OxmlElement(f"w:{side}")
    edge.set(qn("w:val"), "single")
    edge.set(qn("w:sz"), str(size))
    edge.set(qn("w:space"), str(space))
    edge.set(qn("w:color"), color)
    _ppr_child(paragraph, "w:pBdr").append(edge)


def page_left_border(doc, color: str, size: int) -> None:
    borders = OxmlElement("w:pgBorders")
    borders.set(qn("w:offsetFrom"), "page")
    left = OxmlElement("w:left")
    for key, value in (("val", "single"), ("sz", str(size)), ("space", "0"), ("color", color)):
        left.set(qn(f"w:{key}"), value)
    borders.append(left)
    doc.sections[0]._sectPr.append(borders)


def run(paragraph, text: str, size: float, *, bold: bool = False, color: str = "20242A"):
    r = paragraph.add_run(text)
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = RGBColor.from_string(color)
    rfonts = r._element.get_or_add_rPr().get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), FONT)
    rfonts.set(qn("w:hAnsi"), FONT)
    rfonts.set(qn("w:eastAsia"), FONT_EAST_ASIA)
    return r


def para(doc, *, before: float = 0, after: float = 0, line: float = 1.0, keep: bool = False):
    p = doc.add_paragraph()
    fmt = p.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.keep_with_next = keep
    return p


def right_tab(paragraph, position_cm: float) -> None:
    paragraph.paragraph_format.tab_stops.add_tab_stop(Cm(position_cm), 2)  # 2 = right


# ---------- template styles ----------

@dataclass(frozen=True)
class Style:
    accent: str
    rule: str
    margins: tuple[float, float, float, float]  # top, bottom, left, right (cm)
    heading_size: float
    body_size: float
    entry_size: float
    bullet_after: float
    body_line: float
    contact_sep: str

    @property
    def text_width(self) -> float:
        return PAGE_WIDTH_CM - self.margins[2] - self.margins[3]


STYLES = {
    "professional-business": Style("1C4069", "D3DEE8", (1.05, 1.0, 1.7, 1.55), 11.2, 9.75, 10.4, 1.6, 1.16, "  ·  "),
    "technical-project": Style("087F8C", "D7E9E7", (1.05, 1.0, 1.6, 1.6), 11.5, 9.55, 10.3, 1.3, 1.14, "  /  "),
    "one-page-compact": Style("742B32", "742B32", (0.0, 1.0, 1.4, 1.4), 10.8, 9.2, 9.7, 1.8, 1.18, "  ·  "),
    "ats-minimal": Style("111111", "999999", (1.6, 1.5, 1.8, 1.8), 11.0, 10.0, 10.2, 1.6, 1.15, " | "),
}


def new_document(style: Style):
    doc = Document()
    section = doc.sections[0]
    section.start_type = WD_SECTION_START.NEW_PAGE
    section.page_width = Cm(PAGE_WIDTH_CM)
    section.page_height = Cm(29.7)
    top, bottom, left, right = style.margins
    section.top_margin, section.bottom_margin = Cm(top), Cm(bottom)
    section.left_margin, section.right_margin = Cm(left), Cm(right)
    section.header_distance = section.footer_distance = Cm(0.5)
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(style.body_size)
    rfonts = normal.element.get_or_add_rPr().get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), FONT_EAST_ASIA)
    return doc


# ---------- headers ----------

def header_business(doc, style: Style, data: dict) -> None:
    page_left_border(doc, style.accent, 96)
    p = para(doc, after=0.2)
    run(p, data["name"], 22, bold=True, color=style.accent)
    p = para(doc, after=5.5)
    right_tab(p, style.text_width)
    run(p, data.get("target_role", ""), 9.8, bold=True, color=style.accent)
    run(p, "\t" + style.contact_sep.join(data.get("contact", [])), 7.8, color="717B86")
    border(p, "bottom", style.accent, 14)


def header_technical(doc, style: Style, data: dict) -> None:
    p = para(doc, after=2)
    right_tab(p, style.text_width)
    run(p, data["name"], 22, bold=True, color="15333A")
    if data.get("target_role"):
        run(p, "\t" + data["target_role"], 9, bold=True, color=style.accent)
    p = para(doc, after=5)
    p.paragraph_format.left_indent = Cm(0.12)
    run(p, style.contact_sep.join(data.get("contact", [])), 7.9, color="3F6264")
    shade(p, "E6F4F2")


def header_compact(doc, style: Style, data: dict) -> None:
    margin = style.margins[2]
    contact = data.get("contact", [])
    half = (len(contact) + 1) // 2

    def band(height_pt: float):
        p = doc.add_paragraph()
        fmt = p.paragraph_format
        fmt.left_indent = fmt.right_indent = Cm(-margin)
        fmt.space_before = fmt.space_after = Pt(0)
        fmt.line_spacing = Pt(height_pt)
        shade(p, style.accent)
        return p

    run(band(23), " ", 1, color=style.accent)
    p = band(34)
    p.paragraph_format.first_line_indent = Cm(margin)
    right_tab(p, PAGE_WIDTH_CM - margin)
    run(p, data["name"], 23, bold=True, color="FFFFFF")
    run(p, "\t" + style.contact_sep.join(contact[:half]), 7.8, color="F5E7EA")
    p = band(18)
    p.paragraph_format.first_line_indent = Cm(margin)
    right_tab(p, PAGE_WIDTH_CM - margin)
    run(p, data.get("target_role", ""), 8.8, color="F5E7EA")
    run(p, "\t" + style.contact_sep.join(contact[half:]), 7.8, color="F5E7EA")
    run(band(39), " ", 1, color=style.accent)
    para(doc, after=2.5)


def header_plain(doc, style: Style, data: dict) -> None:
    run(para(doc, after=1.5), data["name"], 18, bold=True, color=style.accent)
    if data.get("target_role"):
        run(para(doc, after=1.0), f"{data.get('target_role_label', '求职意向：')}{data['target_role']}", 10, color=style.accent)
    run(para(doc, after=4.0), style.contact_sep.join(data.get("contact", [])), 9.5, color="555555")


HEADERS = {
    "professional-business": header_business,
    "technical-project": header_technical,
    "one-page-compact": header_compact,
    "ats-minimal": header_plain,
}


# ---------- body ----------

def heading(doc, template: str, style: Style, text: str) -> None:
    if template == "professional-business":
        p = para(doc, before=6.5, after=3.3, keep=True)
        run(p, text, style.heading_size, bold=True, color=style.accent)
        border(p, "left", style.accent, 20, 7)
        border(p, "bottom", style.rule, 5)
    elif template == "technical-project":
        p = para(doc, before=6.0, after=3.0, keep=True)
        run(p, text, style.heading_size, bold=True, color=style.accent)
        border(p, "bottom", style.rule, 4)
    elif template == "one-page-compact":
        p = para(doc, before=8.0, after=4.0, keep=True)
        run(p, text, style.heading_size, bold=True, color=style.accent)
        border(p, "bottom", style.rule, 5)
    else:
        p = para(doc, before=7.0, after=3.0, keep=True)
        run(p, text, style.heading_size, bold=True, color=style.accent)
        border(p, "bottom", style.rule, 4)


def summary(doc, template: str, style: Style, text: str) -> None:
    p = para(doc, before=0.8 if template == "professional-business" else 0, after=2.5, line=style.body_line)
    run(p, text, style.body_size + 0.25)
    if template == "professional-business":
        p.paragraph_format.left_indent = p.paragraph_format.right_indent = Cm(0.12)
        shade(p, "EAF1F7")


def skills(doc, template: str, style: Style, items: list[dict]) -> None:
    label_color = "111111" if template == "ats-minimal" else style.accent
    sep = "：" if template == "ats-minimal" else "  "
    for item in items:
        p = para(doc, after=1.2, line=1.08)
        run(p, f"{item['label']}{sep}", style.body_size, bold=True, color=label_color)
        run(p, item["value"], style.body_size)


def entry(doc, style: Style, item: dict) -> None:
    p = para(doc, before=2.5, after=0.2, keep=True)
    right_tab(p, style.text_width)
    run(p, item["org"], style.entry_size + 0.3, bold=True, color="171A1F")
    if item.get("role"):
        run(p, f"  {item['role']}", style.entry_size, bold=True, color=style.accent)
    if item.get("date"):
        run(p, f"\t{item['date']}", style.entry_size - 1, color="626A73")
    for text in item.get("bullets", []):
        bullet(doc, style, text)


def bullet(doc, style: Style, text: str) -> None:
    p = para(doc, after=style.bullet_after, line=1.10)
    p.paragraph_format.left_indent = Cm(0.45)
    p.paragraph_format.first_line_indent = Cm(-0.36)
    run(p, "• ", style.body_size, bold=True, color=style.accent)
    run(p, text, style.body_size)


# ---------- validation, verification and export ----------

def validate(data: dict) -> list[str]:
    errors = []
    if not isinstance(data.get("name"), str) or not data["name"].strip():
        errors.append("name is required")
    if not isinstance(data.get("contact", []), list):
        errors.append("contact must be a list of strings")
    sections = data.get("sections")
    if not isinstance(sections, list) or not sections:
        errors.append("sections must be a non-empty list")
        return errors
    for i, section in enumerate(sections):
        kind = section.get("type")
        if kind not in SECTION_TYPES:
            errors.append(f"sections[{i}].type must be one of {', '.join(SECTION_TYPES)}")
        if not section.get("title"):
            errors.append(f"sections[{i}].title is required")
        if kind == "summary" and not isinstance(section.get("text"), str):
            errors.append(f"sections[{i}].text is required for summary")
        if kind in ("skills", "entries", "bullets") and not isinstance(section.get("items"), list):
            errors.append(f"sections[{i}].items must be a list")
        if kind == "skills":
            for j, item in enumerate(section.get("items", [])):
                if not isinstance(item, dict) or "label" not in item or "value" not in item:
                    errors.append(f"sections[{i}].items[{j}] needs label and value")
        if kind == "entries":
            for j, item in enumerate(section.get("items", [])):
                if not isinstance(item, dict) or not item.get("org"):
                    errors.append(f"sections[{i}].items[{j}].org is required")
    return errors


def header_strings(data: dict) -> list[str]:
    """Header strings; some templates place them side by side, so order is not checked."""
    out = [data["name"]]
    if data.get("target_role"):
        out.append(data["target_role"])
    return out + [c for c in data.get("contact", []) if c]


def expected_strings(data: dict) -> list[str]:
    """Every body string in reading order, used to verify the saved file."""
    out = []
    for section in data["sections"]:
        out.append(section["title"])
        if section["type"] == "summary":
            out.append(section["text"])
        elif section["type"] == "skills":
            for item in section["items"]:
                out += [item["label"], item["value"]]
        elif section["type"] == "entries":
            for item in section["items"]:
                out += [s for s in (item["org"], item.get("role"), item.get("date")) if s]
                out += item.get("bullets", [])
        else:
            out += section["items"]
    return out


def verify(path: Path, data: dict) -> list[str]:
    """Reopen the DOCX and confirm every string is present, in order."""
    text = "\n".join(p.text for p in Document(str(path)).paragraphs)
    problems = [f"missing: {h[:40]}" for h in header_strings(data) if h not in text]
    cursor = 0
    for expected in expected_strings(data):
        found = text.find(expected, cursor)
        if found < 0:
            where = "missing" if expected not in text else "out of order"
            problems.append(f"{where}: {expected[:40]}")
        else:
            cursor = found + len(expected)
    return problems


def find_soffice() -> str | None:
    for candidate in ("soffice", "libreoffice", "/Applications/LibreOffice.app/Contents/MacOS/soffice"):
        found = shutil.which(candidate) or (candidate if Path(candidate).is_file() else None)
        if found:
            return found
    return None


def export_pdf(docx: Path) -> tuple[Path | None, int | None, str | None]:
    soffice = find_soffice()
    if not soffice:
        return None, None, "LibreOffice not found; PDF not created. Export the DOCX to PDF with the host's document tools or Word/WPS."
    with tempfile.TemporaryDirectory() as profile:
        result = subprocess.run(
            [soffice, f"-env:UserInstallation=file://{profile}", "--headless", "--convert-to", "pdf", "--outdir", str(docx.parent), str(docx)],
            capture_output=True, text=True, timeout=180,
        )
    pdf = docx.with_suffix(".pdf")
    if result.returncode != 0 or not pdf.exists():
        return None, None, "PDF conversion failed."
    pages = len(re.findall(rb"/Type\s*/Page(?!s)", pdf.read_bytes()))
    return pdf, pages or None, None


def render(data: dict) -> Document:
    template = LEGACY.get(data.get("template", "ats-minimal"), data.get("template", "ats-minimal"))
    if template not in TEMPLATES:
        raise ValueError(f"unknown template {template!r}; use one of {', '.join(TEMPLATES)}")
    style = STYLES[template]
    doc = new_document(style)
    HEADERS[template](doc, style, data)
    for section in data["sections"]:
        heading(doc, template, style, section["title"])
        kind = section["type"]
        if kind == "summary":
            summary(doc, template, style, section["text"])
        elif kind == "skills":
            skills(doc, template, style, section["items"])
        elif kind == "entries":
            for item in section["items"]:
                entry(doc, style, item)
        else:
            for text in section["items"]:
                bullet(doc, style, text)
    doc.core_properties.title = f"{data['name']} 简历"
    doc.core_properties.author = data["name"]
    return doc


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("input", type=Path, help="resume JSON file")
    parser.add_argument("--out", type=Path, required=True, help="output .docx path")
    parser.add_argument("--template", help="override the template in the JSON")
    parser.add_argument("--pdf", action="store_true", help="also export PDF with LibreOffice when available")
    args = parser.parse_args(argv)

    def fail(message: str, code: int = 1) -> int:
        print(json.dumps({"ok": False, "error": message}, ensure_ascii=False))
        return code

    try:
        data = json.loads(args.input.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        return fail(f"cannot read input JSON: {exc}")
    if args.template:
        data["template"] = args.template
    errors = validate(data)
    if errors:
        return fail("; ".join(errors))
    try:
        doc = render(data)
    except ValueError as exc:
        return fail(str(exc))

    out = args.out if args.out.suffix == ".docx" else args.out.with_suffix(".docx")
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    warnings = verify(out, data)
    pdf = pages = None
    if args.pdf:
        pdf, pages, note = export_pdf(out)
        if note:
            warnings.append(note)
        template = LEGACY.get(data.get("template", ""), data.get("template", ""))
        if pages and template == "one-page-compact" and pages > 1:
            warnings.append(f"one-page-compact rendered to {pages} pages; shorten repeated or low-relevance content")
        if pages and pages > 2:
            warnings.append(f"resume is {pages} pages; Chinese applications usually expect one or two")
    if LEGACY.get(data.get("template", ""), data.get("template", "")) == "one-page-compact" and not pdf:
        warnings.append("one-page-compact puts the name in white text on a coloured band; deliver it as PDF, because some online previewers drop the band and the name becomes invisible")
    print(json.dumps({"ok": not any(w.startswith(("missing", "out of order")) for w in warnings),
                      "docx": str(out), "pdf": str(pdf) if pdf else None, "pages": pages,
                      "warnings": warnings}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())
