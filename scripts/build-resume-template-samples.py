from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "skills" / "job-fit-assistant" / "assets" / "resume-templates"
FONT = "Arial"
FONT_EAST_ASIA = "Microsoft YaHei"

PROFILE = (
    "3 年互联网产品运营经验，聚焦效率工具的新用户激活、功能采用与数据分析。"
    "能够使用 SQL 和 Looker Studio 拆解注册、激活与留存漏斗，设计并复盘 onboarding 实验，"
    "协同产品、设计和研发推动需求上线。"
)

SKILLS = [
    ("数据分析", "SQL、Excel、Looker Studio、漏斗分析"),
    ("用户运营", "onboarding 优化、A/B 测试、用户访谈、社群运营"),
    ("产品协作", "需求文档、埋点验收、跨团队推进、上线复盘"),
    ("AI 与英语", "大模型回答质量评测、英文产品文档阅读"),
]

DISPLAY_SKILLS = [
    ("数据分析", "SQL、Excel、Looker Studio、漏斗分析"),
    ("用户运营", "onboarding 优化、A/B 测试、用户访谈、社群运营"),
    ("产品协作", "需求文档、埋点验收、跨团队推进、AI 回答质量评测"),
]

EXPERIENCES = [
    (
        "云笺科技",
        "产品运营",
        "2023 年 3 月至今｜杭州",
        [
            "负责个人知识管理产品的新用户激活，结合访谈和行为漏斗定位首日流失节点；推动 4 轮 onboarding 实验，次日核心功能激活率由 41% 提升至 55%。",
            "使用 SQL 和 Looker Studio 搭建注册、激活、留存周报，支持产品团队每周复盘；将人工整理时间从约 3 小时缩短到 30 分钟。",
            "联动产品、设计和研发上线模板推荐与新手任务，负责需求说明、埋点验收和上线复盘。",
            "运营用户社群并整理高频问题，每月输出产品反馈报告，帮助团队确定下一版本的体验优化优先级。",
        ],
    ),
    (
        "星途教育",
        "用户运营实习生",
        "2022 年 6 月至 2022 年 12 月｜远程",
        [
            "维护课程用户社群，整理咨询与完课数据，协助优化开课提醒和学习路径。",
            "制作基础数据看板，跟踪报名、到课和完课情况。",
        ],
    ),
]

DISPLAY_EXPERIENCES = [
    (
        "云笺科技",
        "产品运营",
        "2023.03—至今",
        [
            "结合访谈和行为漏斗定位首日流失节点，推动 4 轮 onboarding 实验，次日核心功能激活率由 41% 提升至 55%。",
            "使用 SQL 和 Looker Studio 搭建注册、激活、留存周报，将人工整理时间从约 3 小时缩短到 30 分钟。",
            "联动产品、设计和研发上线模板推荐与新手任务，负责需求说明、埋点验收与复盘，并整理社群高频问题形成反馈报告。",
        ],
    ),
    (
        "星途教育",
        "用户运营实习生",
        "2022.06—2022.12",
        [
            "维护课程用户社群，整理咨询与完课数据，协助优化开课提醒和学习路径。",
            "制作基础数据看板，跟踪报名、到课和完课情况。",
        ],
    ),
]

PROJECT = (
    "LLM 回答质量评测实验",
    "个人实验｜2025 年 11 月至 2026 年 1 月",
    [
        "为 120 条公开问答样本设计正确性、完整性和可引用性三维评分规则。",
        "比较 3 个通用模型的输出并复核分歧样本，形成评测记录和问题分类。",
        "项目未用于生产环境，也没有商业客户。",
    ],
)


def set_cell_shading(paragraph, fill: str) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    shading = ppr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        ppr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_bottom_border(paragraph, color: str, size: int = 8) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    borders = ppr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        ppr.append(borders)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)


def set_left_border(paragraph, color: str, size: int = 18, space: int = 6) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    borders = ppr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        ppr.append(borders)
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), str(size))
    left.set(qn("w:space"), str(space))
    left.set(qn("w:color"), color)
    borders.append(left)


def set_page_left_border(doc: Document, color: str, size: int = 30) -> None:
    sect_pr = doc.sections[0]._sectPr
    borders = OxmlElement("w:pgBorders")
    borders.set(qn("w:offsetFrom"), "page")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), str(size))
    left.set(qn("w:space"), "0")
    left.set(qn("w:color"), color)
    borders.append(left)
    sect_pr.append(borders)


def set_font(run, size: float, *, bold: bool = False, color: str = "20242A") -> None:
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for key in ("w:ascii", "w:hAnsi"):
        rfonts.set(qn(key), FONT)
    rfonts.set(qn("w:eastAsia"), FONT_EAST_ASIA)


def spacing(paragraph, before: float = 0, after: float = 0, line: float = 1.0) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def setup_doc(*, top: float, bottom: float, left: float, right: float) -> Document:
    doc = Document()
    section = doc.sections[0]
    section.start_type = WD_SECTION_START.NEW_PAGE
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(top)
    section.bottom_margin = Cm(bottom)
    section.left_margin = Cm(left)
    section.right_margin = Cm(right)
    section.header_distance = Cm(0.5)
    section.footer_distance = Cm(0.5)
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10.2)
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_EAST_ASIA)
    return doc


def add_text(doc: Document, text: str, *, size: float, bold: bool = False,
             color: str = "20242A", align=None, before: float = 0,
             after: float = 0, line: float = 1.0, keep: bool = False):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    spacing(p, before, after, line)
    p.paragraph_format.keep_with_next = keep
    set_font(p.add_run(text), size, bold=bold, color=color)
    return p


def add_bullet(doc: Document, text: str, *, size: float, after: float,
               color: str = "20242A", left_cm: float = 0.45):
    p = doc.add_paragraph()
    spacing(p, 0, after, 1.10)
    p.paragraph_format.left_indent = Cm(left_cm)
    p.paragraph_format.first_line_indent = Cm(-0.36)
    set_font(p.add_run("• "), size, bold=True, color=color)
    set_font(p.add_run(text), size, color="20242A")
    return p


def add_role(doc: Document, company: str, role: str, date: str, *, accent: str,
             size: float, compact: bool = False):
    p = doc.add_paragraph()
    spacing(p, 2.5 if not compact else 1.2, 0.2, 1.0)
    p.paragraph_format.keep_with_next = True
    set_font(p.add_run(company), size + 0.3, bold=True, color="171A1F")
    set_font(p.add_run(f"  {role}"), size, bold=True, color=accent)
    tab = p.paragraph_format.tab_stops
    tab.add_tab_stop(Cm(16.2), 2)
    set_font(p.add_run(f"\t{date}"), size - 1, color="626A73")


def business() -> Path:
    navy = "1C4069"
    pale = "EAF1F7"
    doc = setup_doc(top=1.05, bottom=1.0, left=1.7, right=1.55)
    set_page_left_border(doc, navy, 96)
    p = add_text(doc, "", size=10, after=0.2, line=1.0)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(15.9), 2)
    set_font(p.add_run("林一然"), 22, bold=True, color=navy)
    p = add_text(doc, "", size=9, after=5.5, line=1.0)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(15.9), 2)
    set_font(p.add_run("产品运营"), 9.8, bold=True, color=navy)
    set_font(p.add_run("\t杭州  ·  lin.yiran@example.com  ·  138 0000 0000"), 7.8, color="717B86")
    set_bottom_border(p, navy, 14)

    def heading(text: str):
        p = add_text(doc, text, size=11.2, bold=True, color=navy, before=6.5, after=3.3, keep=True)
        set_left_border(p, navy, 20, 7)
        set_bottom_border(p, "D3DEE8", 5)

    heading("个人概况")
    p = add_text(doc, PROFILE, size=10.0, before=0.8, after=3.0, line=1.16)
    p.paragraph_format.left_indent = Cm(0.12)
    p.paragraph_format.right_indent = Cm(0.12)
    set_cell_shading(p, pale)
    heading("核心技能")
    for label, value in DISPLAY_SKILLS:
        p = add_text(doc, "", size=9.8, after=1.1)
        set_font(p.add_run(f"{label}  "), 9.8, bold=True, color=navy)
        set_font(p.add_run(value), 9.8)
    heading("工作经历")
    for company, role, date, bullets in DISPLAY_EXPERIENCES:
        add_role(doc, company, role, date, accent=navy, size=10.4)
        for item in bullets:
            add_bullet(doc, item, size=9.75, after=1.6, color=navy)
    heading("个人项目")
    add_role(doc, "大模型回答质量评测实验", "个人实验", "2025.11—2026.01", accent=navy, size=10.4)
    for item in PROJECT[2][:2]:
        add_bullet(doc, item, size=9.75, after=1.4, color=navy)
    heading("教育经历")
    add_role(doc, "华东理工大学", "工商管理 本科", "2018.09—2022.06", accent=navy, size=10.4)
    doc.core_properties.title = "林一然 专业商务版简历"
    path = OUTPUT / "professional-business.docx"
    doc.save(path)
    return path


def technical() -> Path:
    teal = "087F8C"
    dark = "15333A"
    pale = "E6F4F2"
    doc = setup_doc(top=1.05, bottom=1.0, left=1.6, right=1.6)
    p = add_text(doc, "", size=10, after=2, line=1.0)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(16.0), 2)
    set_font(p.add_run("林一然"), 22, bold=True, color=dark)
    set_font(p.add_run("\t产品运营"), 7.4, bold=True, color=teal)
    p = add_text(doc, "杭州  /  lin.yiran@example.com  /  138 0000 0000", size=7.9,
                 color="3F6264", after=5, line=1.0)
    p.paragraph_format.left_indent = Cm(0.12)
    set_cell_shading(p, pale)

    def heading(text: str):
        p = doc.add_paragraph()
        spacing(p, 6.0, 3.0, 1.0)
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.tab_stops.add_tab_stop(Cm(16.0), 2)
        set_font(p.add_run(text), 11.5, bold=True, color=teal)
        set_bottom_border(p, "D7E9E7", 4)

    heading("个人概况")
    add_text(doc, PROFILE, size=9.8, after=1.5, line=1.14)
    heading("核心技能")
    for label, value in DISPLAY_SKILLS:
        p = add_text(doc, "", size=9.7, after=1.0)
        set_font(p.add_run(f"{label}  "), 9.55, bold=True, color=teal)
        set_font(p.add_run(value), 9.7)
    heading("个人项目")
    add_role(doc, "大模型回答质量评测实验", "个人实验", "2025.11—2026.01", accent=teal, size=10.3)
    for item in PROJECT[2][:2]:
        add_bullet(doc, item, size=9.55, after=1.3, color=teal)
    heading("工作经历")
    for company, role, date, bullets in DISPLAY_EXPERIENCES:
        add_role(doc, company, role, date, accent=teal, size=10.3)
        for item in bullets:
            add_bullet(doc, item, size=9.55, after=1.3, color=teal)
    heading("教育经历")
    add_role(doc, "华东理工大学", "工商管理 本科", "2018.09—2022.06", accent=teal, size=10.3)
    doc.core_properties.title = "林一然 技术项目版简历"
    path = OUTPUT / "technical-project.docx"
    doc.save(path)
    return path


def compact() -> Path:
    wine = "742B32"
    doc = setup_doc(top=0, bottom=1.0, left=1.4, right=1.4)
    def band_line(height_pt: float):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(-1.4)
        p.paragraph_format.right_indent = Cm(-1.4)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = Pt(height_pt)
        set_cell_shading(p, wine)
        return p

    p = band_line(23)
    set_font(p.add_run(" "), 1, color=wine)
    p = band_line(34)
    p.paragraph_format.first_line_indent = Cm(1.4)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(18.2), 2)
    set_font(p.add_run("林一然"), 23, bold=True, color="FFFFFF")
    set_font(p.add_run("\t杭州  ·  138 0000 0000"), 7.8, color="F5E7EA")
    p = band_line(18)
    p.paragraph_format.first_line_indent = Cm(1.4)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(18.2), 2)
    set_font(p.add_run("产品运营"), 8.8, color="F5E7EA")
    set_font(p.add_run("\tlin.yiran@example.com"), 7.8, color="F5E7EA")
    p = band_line(39)
    set_font(p.add_run(" "), 1, color=wine)
    spacer = doc.add_paragraph()
    spacing(spacer, 0, 2.5, 1.0)

    def heading(text: str):
        p = add_text(doc, text, size=10.8, bold=True, color=wine, before=8.0, after=4.0, keep=True)
        set_bottom_border(p, wine, 5)

    heading("个人概况")
    add_text(doc, PROFILE, size=9.4, after=2.0, line=1.18)
    heading("核心技能")
    for label, value in DISPLAY_SKILLS:
        p = add_text(doc, "", size=9.1, after=1.5, line=1.08)
        set_font(p.add_run(f"{label}  "), 9.1, bold=True, color=wine)
        set_font(p.add_run(value), 9.1)
    heading("工作经历")
    for company, role, date, bullets in DISPLAY_EXPERIENCES:
        add_role(doc, company, role, date, accent=wine, size=9.7, compact=True)
        for item in bullets:
            add_bullet(doc, item, size=9.2, after=1.8, color=wine, left_cm=0.34)
    heading("个人项目")
    add_role(doc, "大模型回答质量评测实验", "个人实验", "2025.11—2026.01", accent=wine, size=9.7, compact=True)
    for item in PROJECT[2][:2]:
        add_bullet(doc, item, size=9.2, after=1.7, color=wine, left_cm=0.34)
    heading("教育经历")
    add_role(doc, "华东理工大学", "工商管理 本科", "2018.09—2022.06", accent=wine, size=9.7, compact=True)
    doc.core_properties.title = "林一然 一页紧凑版简历"
    path = OUTPUT / "one-page-compact.docx"
    doc.save(path)
    return path


def plain() -> Path:
    black = "111111"
    gray = "555555"
    doc = setup_doc(top=1.6, bottom=1.5, left=1.8, right=1.8)
    add_text(doc, "林一然", size=18, bold=True, color=black, after=1.5)
    add_text(doc, "求职意向：产品运营", size=10, color=black, after=1.0)
    add_text(doc, "杭州 | lin.yiran@example.com | 138 0000 0000", size=9.5, color=gray, after=4.0)

    def heading(text: str):
        p = add_text(doc, text, size=11, bold=True, color=black, before=7.0, after=3.0, keep=True)
        set_bottom_border(p, "999999", 4)

    heading("个人概况")
    add_text(doc, PROFILE, size=10, after=2.0, line=1.15)
    heading("核心技能")
    for label, value in DISPLAY_SKILLS:
        p = add_text(doc, "", size=10, after=1.2)
        set_font(p.add_run(f"{label}："), 10, bold=True, color=black)
        set_font(p.add_run(value), 10, color=black)
    heading("工作经历")
    for company, role, date, bullets in DISPLAY_EXPERIENCES:
        add_role(doc, company, role, date, accent=black, size=10.2)
        for item in bullets:
            add_bullet(doc, item, size=10, after=1.6, color=black)
    heading("个人项目")
    add_role(doc, "大模型回答质量评测实验", "个人实验", "2025.11—2026.01", accent=black, size=10.2)
    for item in PROJECT[2][:2]:
        add_bullet(doc, item, size=10, after=1.4, color=black)
    heading("教育经历")
    add_role(doc, "华东理工大学", "工商管理 本科", "2018.09—2022.06", accent=black, size=10.2)
    doc.core_properties.title = "林一然 极简黑白版简历"
    path = OUTPUT / "ats-minimal.docx"
    doc.save(path)
    return path


if __name__ == "__main__":
    OUTPUT.mkdir(parents=True, exist_ok=True)
    for output in (business(), technical(), compact(), plain()):
        print(output)
